"""
One-time / on-demand ingest of the ONS Price Index of Private Rents (PIPR).

Unlike the Land Registry feed, ONS has no stable "latest" URL for this dataset —
the release date is baked into both the landing page path and the file URL, and
ONS's old dataset API is fully decommissioned. So this scrapes the dataset
landing page for every edition's download link + publish date, and picks the
newest one.

Usage:
  python ingest/load_rents.py
"""

import csv
import os
import re
import sys
from datetime import datetime

import boto3
import duckdb
import requests
from dotenv import load_dotenv
from openpyxl import load_workbook

load_dotenv()

BUCKET = os.getenv("AWS_BUCKET", "victor-london-housing-analytics")
REGION = os.getenv("AWS_REGION", "eu-west-2")

DATASET_PAGE = (
    "https://www.ons.gov.uk/economy/inflationandpriceindices/datasets/"
    "priceindexofprivaterentsukmonthlypricestatistics"
)

# England (unitary/non-metropolitan district/metropolitan district/London
# borough) and Wales (unitary authority) local authority area codes — matches
# the same England & Wales scope Land Registry already covers.
LA_CODE_PREFIXES = ("E06", "E07", "E08", "E09", "W06")

DOWNLOAD_LINK_RE = re.compile(
    r'href="(?P<href>/file\?uri=[^"]+\.xlsx)"[^>]*'
    r'aria-label="Download[^"]*:\s*(?P<date>\d{1,2} \w+ \d{4})\s+in xlsx format"'
)


def find_latest_xlsx_url() -> tuple[datetime, str]:
    print("Finding latest PIPR release...")
    resp = requests.get(DATASET_PAGE, timeout=30)
    resp.raise_for_status()

    candidates = []
    for m in DOWNLOAD_LINK_RE.finditer(resp.text):
        try:
            dt = datetime.strptime(m.group("date"), "%d %B %Y")
        except ValueError:
            continue
        candidates.append((dt, "https://www.ons.gov.uk" + m.group("href")))

    if not candidates:
        raise RuntimeError("No PIPR download links found on the ONS dataset page — page layout may have changed")

    candidates.sort(key=lambda c: c[0])
    latest_date, latest_url = candidates[-1]
    print(f"  ✓ latest edition: {latest_date:%d %B %Y}")
    return latest_date, latest_url


def download(url: str, dest: str) -> None:
    print("Downloading from ONS...")
    r = requests.get(url, timeout=120)
    r.raise_for_status()
    with open(dest, "wb") as f:
        f.write(r.content)
    print(f"  ✓ {len(r.content) / 1_000_000:.1f} MB downloaded")


def parse_and_filter(xlsx_path: str, csv_path: str) -> int:
    print("Parsing Table 1 and filtering to local authorities...")
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb["Table 1"]

    header = None
    rows_written = 0
    with open(csv_path, "w", newline="") as f:
        writer = csv.writer(f)
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 2:
                header = row
                writer.writerow(header)
                continue
            if i < 3:
                continue
            area_code = row[1]
            if not area_code or not str(area_code).startswith(LA_CODE_PREFIXES):
                continue
            writer.writerow(row)
            rows_written += 1

    wb.close()
    print(f"  ✓ {rows_written:,} local-authority rows kept")
    return rows_written


def upload_to_s3(local_path: str, s3_key: str) -> str:
    print(f"Uploading to s3://{BUCKET}/{s3_key} ...")
    client = boto3.client(
        "s3",
        aws_access_key_id=os.getenv("AWS_KEY_ID"),
        aws_secret_access_key=os.getenv("AWS_SECRET"),
        region_name=REGION,
    )
    client.upload_file(local_path, BUCKET, s3_key)
    s3_uri = f"s3://{BUCKET}/{s3_key}"
    print(f"  ✓ uploaded to {s3_uri}")
    return s3_uri


def load_into_duckdb(s3_uri: str, db_path: str) -> int:
    print("Loading into DuckDB...")
    con = duckdb.connect(db_path)
    try:
        con.execute(f"""
            INSTALL httpfs;
            LOAD httpfs;
            CREATE SECRET IF NOT EXISTS aws_secret (
                TYPE S3,
                KEY_ID '{os.getenv("AWS_KEY_ID")}',
                SECRET '{os.getenv("AWS_SECRET")}',
                REGION '{REGION}'
            );
        """)
        con.execute(f"""
            CREATE OR REPLACE TABLE raw_rents AS
            SELECT
                *,
                '{s3_uri}' AS _source_file,
                current_timestamp AS _loaded_at
            FROM read_csv_auto('{s3_uri}', header=true)
        """)
        count = con.execute("SELECT COUNT(*) FROM raw_rents").fetchone()[0]
        print(f"  ✓ raw_rents: {count:,} rows")
        return count
    finally:
        con.close()


if __name__ == "__main__":
    db_path = os.path.join(
        os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "london_housing.duckdb"
    )
    latest_date, url = find_latest_xlsx_url()
    tmp_xlsx = "/tmp/pipr-latest.xlsx"
    tmp_csv = f"/tmp/pipr-{latest_date:%Y-%m}.csv"

    try:
        download(url, tmp_xlsx)
        row_count = parse_and_filter(tmp_xlsx, tmp_csv)
        if row_count == 0:
            print("No local-authority rows found — aborting without writing to S3 or DuckDB.")
            sys.exit(1)
        s3_key = f"rents/pipr-{latest_date:%Y-%m}.csv"
        s3_uri = upload_to_s3(tmp_csv, s3_key)
        load_into_duckdb(s3_uri, db_path)
        print("\nDone.")
    finally:
        for tmp in (tmp_xlsx, tmp_csv):
            if os.path.exists(tmp):
                os.remove(tmp)
